#!/usr/bin/env python3
"""Build the editable Orange Romania Delivery Pitch deck (3 slides, Northgate light system)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

ROOT = "/Users/imgauravraj/Claude Code/Claude Agents/My Company Website/northgate"
LOGO = os.path.join(ROOT, "public/brand/northgate-logo-navy-transparent.png")
OUT  = os.path.join(ROOT, "Orange Romania/Orange-Romania-Delivery-Pitch.pptx")

INK    = RGBColor(0x0C, 0x19, 0x45)
INK2   = RGBColor(0x45, 0x4F, 0x74)
INK3   = RGBColor(0x71, 0x79, 0x9C)
ACCENT = RGBColor(0x2F, 0x4B, 0xD0)
CYAN   = RGBColor(0x05, 0xAF, 0xF2)
PERI   = RGBColor(0x5B, 0x73, 0xE6)
WASH   = RGBColor(0xEE, 0xF5, 0xFE)
LINE   = RGBColor(0xD7, 0xE0, 0xF1)

HEAD = "Lexend"
BODY = "Source Sans 3"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

L = 0.65
CW = 13.333 - 2 * L
GAP = 0.55
COLW = (CW - 2 * GAP) / 3
COLX = [L, L + COLW + GAP, L + 2 * (COLW + GAP)]


def no_shadow(shape):
    el = shape._element
    # drop the theme style reference (carries the preset outer shadow)
    for st in el.findall(qn('p:style')):
        el.remove(st)
    spPr = shape._element.spPr
    for tag in ('a:effectLst', 'a:effectDag'):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    spPr.append(etree.SubElement(spPr, qn('a:effectLst')))


def _spc(run, pts):
    run.font._rPr.set('spc', str(int(pts * 100)))


def text(slide, x, y, w, h, runs, *, size, color, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, leading=1.12, spc=0.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, bold)]
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = leading
    for t, b in runs:
        r = p.add_run()
        r.text = t
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.color.rgb = color
        if spc:
            _spc(r, spc)
    return tb


def rect(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    no_shadow(s)
    return s


def rounded(slide, x, y, w, h, fill, line=LINE, radius=0.055):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        s.adjustments[0] = radius
    except Exception:
        pass
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line; s.line.width = Pt(0.75)
    s.shadow.inherit = False
    no_shadow(s)
    return s


def chevron(slide, x, y, size, color):
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(size * 1.15), Inches(size))
    s.fill.background()
    s.line.color.rgb = color; s.line.width = Pt(1.4)
    s.shadow.inherit = False
    no_shadow(s)
    return s


def base(slide, eyebrow, title, lede_runs, *, lede_y=2.5):
    slide.shapes.add_picture(LOGO, Inches(L), Inches(0.44), height=Inches(0.36))
    text(slide, 13.333 - L - 4.4, 0.5, 4.4, 0.3, "CONFIDENTIAL   ·   DISCUSSION DRAFT",
         size=8, color=INK3, font=HEAD, align=PP_ALIGN.RIGHT, spc=1.6)
    text(slide, L, 1.0, CW, 0.32, eyebrow, size=10.5, color=ACCENT, font=HEAD, bold=True, spc=2.6)
    text(slide, L, 1.36, CW * 0.92, 1.5, title, size=27, color=INK, font=HEAD, bold=True, leading=1.05)
    text(slide, L, lede_y, CW * 0.78, 1.1, lede_runs, size=13.5, color=INK2, leading=1.4)


def col_header(slide, i, kicker, h2, *, y, rule_color=CYAN):
    x = COLX[i]
    rect(slide, x, y, COLW, 0.03, rule_color)
    if i < 2:
        chevron(slide, x + COLW + GAP / 2 - 0.1, y - 0.1, 0.2, rule_color)
    text(slide, x, y + 0.15, COLW, 0.28, kicker, size=8, color=INK3, font=HEAD, bold=True, spc=1.6)
    text(slide, x, y + 0.44, COLW, 0.8, h2, size=16, color=INK, font=HEAD, bold=True, leading=1.1)


def bullets(slide, i, items, *, y, size=11, after=7):
    x = COLX[i]
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(COLW), Inches(2.0))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for n, it in enumerate(items):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.line_spacing = 1.26; p.space_after = Pt(after)
        r = p.add_run(); r.text = "—  "
        r.font.name = HEAD; r.font.size = Pt(size - 0.5); r.font.color.rgb = PERI; r.font.bold = True
        r2 = p.add_run(); r2.text = it
        r2.font.name = BODY; r2.font.size = Pt(size); r2.font.color.rgb = INK2


def band(slide, y, glyph, kicker, title, sub, *, metric=None):
    h = 1.0
    rounded(slide, L, y, CW, h, WASH)
    rect(slide, L, y + 0.05, 0.05, h - 0.1, ACCENT)
    text(slide, L + 0.3, y, 0.85, h, glyph, size=25, color=ACCENT, font=HEAD,
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    tw = CW - 1.5 - (2.7 if metric else 0.4)
    text(slide, L + 1.15, y + 0.13, tw, 0.26, kicker, size=8, color=ACCENT, font=HEAD, bold=True, spc=1.9)
    text(slide, L + 1.15, y + 0.36, tw, 0.4, title, size=14.5, color=INK, font=HEAD, bold=True, leading=1.1)
    text(slide, L + 1.15, y + 0.66, tw, 0.3, sub, size=10, color=INK2, leading=1.28)
    if metric:
        mx = L + CW - 2.45
        rect(slide, mx - 0.3, y + 0.22, 0.014, h - 0.44, LINE)
        text(slide, mx, y + 0.15, 2.3, 0.5, metric[0], size=20, color=ACCENT, font=HEAD, bold=True, align=PP_ALIGN.RIGHT)
        text(slide, mx, y + 0.64, 2.3, 0.3, metric[1], size=8, color=INK3, font=HEAD, align=PP_ALIGN.RIGHT, spc=1.1)


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ══ SLIDE 1 ══════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
base(s1, "THE OPPORTUNITY   —   MY READ",
     "What TCR wants to build with Orange Romania",
     [("Orange runs a pilot of a ", False),
      ("national-style A2P sender & campaign registry", True),
      (". TCR supplies the identity platform and stays ", False),
      ("out of the message path", True),
      ("; Orange supplies the market leverage to make senders register.", False)])
CY1 = 3.62
col_header(s1, 0, "1  ·  REGISTER", "Who & what", y=CY1)
text(s1, COLX[0], CY1 + 1.0, COLW, 1.3,
     "Brands, CSPs and aggregators register their identity and every campaign’s use case in TCR.",
     size=11, color=INK2, leading=1.38)
col_header(s1, 1, "2  ·  VERIFY & SHARE", "Trusted sender data", y=CY1)
text(s1, COLX[1], CY1 + 1.0, COLW, 1.3,
     "TCR verifies identity by role and feeds sender & campaign data to Orange — not the traffic.",
     size=11, color=INK2, leading=1.38)
col_header(s1, 2, "3  ·  ACT", "Manage the traffic", y=CY1)
text(s1, COLX[2], CY1 + 1.0, COLW, 1.3,
     "Orange’s firewall acts on registration status: curb bad senders, bill and prioritise verified ones.",
     size=11, color=INK2, leading=1.38)
band(s1, 5.95, "✓", "WHY IT MATTERS TO ORANGE",
     "Smishing & fraud down   ·   grey-route A2P recaptured as billable traffic",
     "Plus the foundation for verified-sender and RCS business-messaging revenue.")
text(s1, L, 7.06, CW, 0.3,
     [("Their ask   —   ", False),
      ("technical discovery → pilot scope → integration & training → ANCOM alignment", True)],
     size=9, color=INK3, font=HEAD, spc=0.3)
notes(s1, "My read of the deck. TCR is proposing Orange Romania anchor a national-style A2P sender "
          "& campaign registry: register WHO + WHAT, TCR verifies identity by role and shares "
          "sender data with Orange (never the traffic), Orange's firewall acts on it. Prize for "
          "Orange: fraud/smishing down, grey-route A2P converted to billable traffic, base for "
          "verified-sender and RCS revenue. Their four next steps are on the footer.")

# ══ SLIDE 2 ══════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
base(s2, "HOW I CONTRIBUTE",
     "From solution design to go-live — I own that path",
     [("20 years running operator delivery programmes — ", False),
      ("first technical workshop to live service", True), (".", False)])
CY2 = 3.62
col_header(s2, 0, "STEP 01", "Drive technical workshops", y=CY2)
bullets(s2, 0, ["End-to-end solution design", "Integration & data flows",
                "Solution signed off both sides"], y=CY2 + 1.0)
col_header(s2, 1, "STEP 02", "Define the pilot scope", y=CY2)
bullets(s2, 1, ["Features in / out", "Acceptance criteria", "Business KPIs"], y=CY2 + 1.0)
col_header(s2, 2, "STEP 03", "Deliver to go-live", y=CY2)
bullets(s2, 2, ["Sprint plan & execution", "Project plan & stakeholders",
                "UAT & go-live management"], y=CY2 + 1.0)
band(s2, 5.95, "⇄", "RUNS ACROSS EVERY STEP",
     "The bridge — Orange business   ⇄   TCR technical",
     "One person accountable for translating requirements, constraints and decisions — both directions.")
notes(s2, "How I contribute, mapped to their next steps. 1) Drive the technical workshops to a "
          "signed end-to-end solution. 2) Define the pilot scope - features, acceptance criteria, "
          "business KPIs. 3) Deliver to go-live - sprint plan & execution, project plan, "
          "stakeholder management, UAT, go-live. Underpinning all of it: I am the bridge between "
          "Orange's business teams and TCR's technical teams, translating both directions.")

# ══ SLIDE 3 ══════════════════════════════════════════════
s3 = prs.slides.add_slide(BLANK)
base(s3, "CASE STUDY   —   TIER-1 EUROPEAN MOBILE OPERATOR",
     "Turning grey-route A2P into billable traffic",
     [("A large share of enterprise SMS was arriving via grey routes and unregistered "
       "senders — ", False),
      ("no line of sight on the sender, and A2P revenue leaking", True), (".", False)])
CY3 = 3.5
cols3 = [
    ("THE CHALLENGE", "No line of sight on A2P",
     ["Grey-route & unregistered traffic", "Sender unknown per message",
      "Leaking revenue, rising smishing"], LINE, INK3),
    ("WHAT WE BUILT", "Registry + firewall policy",
     ["Central sender & campaign registry", "Wired into the messaging firewall",
      "Violation-metrics feedback loop"], LINE, INK3),
    ("MY ROLE — SOLUTION & SCOPE LEAD", "Workshop to go-live",
     ["Ran workshops to a signed solution", "Owned scope: features, criteria, KPIs",
      "Bridged business & vendor tech; UAT + go-live"], ACCENT, ACCENT),
]
for i, (kick, h3, its, rc, kc) in enumerate(cols3):
    x = COLX[i]
    rect(s3, x, CY3, COLW, 0.03, rc)
    text(s3, x, CY3 + 0.15, COLW, 0.28, kick, size=8, color=kc, font=HEAD, bold=True, spc=1.4)
    text(s3, x, CY3 + 0.43, COLW, 0.6, h3, size=14.5, color=INK, font=HEAD, bold=True, leading=1.1)
    bullets(s3, i, its, y=CY3 + 1.06, size=10.5, after=6)
band(s3, 5.95, "↗", "OUTCOME",
     "Grey-route A2P brought under registration and billing",
     "Traffic converted from grey to registered, billable routes — and a repeatable model for the next markets.",
     metric=("[ +__% ]", "BILLABLE A2P VOLUME"))
notes(s3, "Case study - anonymised Tier-1 European mobile operator. Same class of solution as TCR "
          "proposes. Challenge: grey-route / unregistered A2P, no sender visibility, revenue "
          "leaking. What we built: central sender & campaign registry + firewall policy "
          "integration + violation feedback loop. My role (Solution & Scope Lead): ran workshops "
          "to a signed solution, owned pilot scope (features / acceptance criteria / business "
          "KPIs), bridged operator business and vendor technical teams, drove UAT and go-live. "
          "Outcome: grey-route A2P brought under registration and billing. "
          "ACTION: replace [ +__% ] with the real figure you can defend, or swap for a "
          "qualitative marker (e.g. 'grey -> registered, rolled to 2 further markets').")

prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT), "bytes")
